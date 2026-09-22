"""
Скрипт создания красивой Excel-таблицы ТМ и TR дисков Cobblemon
на основе JSON-файлов сканирования бочек.
Вкладки: 'TM-disks' и 'TR-disks'
Колонки: ID | Номер | Название | Тип | Количество
"""

import os
import glob
import json
import re
import zipfile
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Цвета типов покемонов (HEX) и перевод
TYPE_DATA = {
    'normal':   {'ru': 'Обычный',     'bg': 'A8A878', 'fg': 'FFFFFF'},
    'fire':     {'ru': 'Огненный',    'bg': 'F08030', 'fg': 'FFFFFF'},
    'water':    {'ru': 'Водный',      'bg': '6890F0', 'fg': 'FFFFFF'},
    'grass':    {'ru': 'Травяной',    'bg': '78C850', 'fg': 'FFFFFF'},
    'electric': {'ru': 'Электрический','bg': 'F8D030', 'fg': '000000'},
    'ice':      {'ru': 'Ледяной',     'bg': '98D8D8', 'fg': '000000'},
    'fighting': {'ru': 'Боевой',      'bg': 'C03028', 'fg': 'FFFFFF'},
    'poison':   {'ru': 'Ядовитый',    'bg': 'A040A0', 'fg': 'FFFFFF'},
    'ground':   {'ru': 'Земляной',    'bg': 'E0C068', 'fg': '000000'},
    'flying':   {'ru': 'Летающий',    'bg': 'A890F0', 'fg': 'FFFFFF'},
    'psychic':  {'ru': 'Психический', 'bg': 'F85888', 'fg': 'FFFFFF'},
    'bug':      {'ru': 'Насекомое',   'bg': 'A8B820', 'fg': 'FFFFFF'},
    'rock':     {'ru': 'Каменный',    'bg': 'B8A038', 'fg': 'FFFFFF'},
    'ghost':    {'ru': 'Призрачный',  'bg': '705898', 'fg': 'FFFFFF'},
    'dragon':   {'ru': 'Драконий',    'bg': '7038F8', 'fg': 'FFFFFF'},
    'steel':    {'ru': 'Стальной',    'bg': 'B8B8D0', 'fg': '000000'},
    'dark':     {'ru': 'Тёмный',      'bg': '705848', 'fg': 'FFFFFF'},
    'fairy':    {'ru': 'Волшебный',   'bg': 'EE99AC', 'fg': '000000'},
}

def load_move_metadata():
    """Загрузка базы данных атак из JAR модов."""
    moves_meta = {}
    
    # 1. Загрузка из SimpleTMs jar
    simpletm_jars = glob.glob(r'G:\curseforge\minecraft\Instances\Society Sunlit Cobblemon\mods\*SimpleTMs*.jar')
    if simpletm_jars:
        with zipfile.ZipFile(simpletm_jars[0], 'r') as z:
            if 'moves.json' in z.namelist():
                raw_moves = json.loads(z.read('moves.json').decode('utf-8'))
                for m in raw_moves:
                    m_key = m.get('move', '').lower().replace(' ', '').replace('-', '')
                    moves_meta[m_key] = {
                        'formatname': m.get('formatname', ''),
                        'typing': m.get('typing', 'normal').lower(),
                        'id': m.get('id', '')
                    }

    # 2. Загрузка русских названий атак из Cobblemon
    ru_names = {}
    for jar_path in glob.glob(r'G:\curseforge\minecraft\Instances\Society Sunlit Cobblemon\mods\*.jar'):
        try:
            with zipfile.ZipFile(jar_path, 'r') as z:
                for n in z.namelist():
                    if 'ru_ru.json' in n:
                        try:
                            d = json.loads(z.read(n).decode('utf-8'))
                            for k, v in d.items():
                                if 'cobblemon.move.' in k and not k.endswith('.desc'):
                                    m_id = k.replace('cobblemon.move.', '').lower().replace(' ', '').replace('-', '')
                                    ru_names[m_id] = v
                        except:
                            pass
        except:
            pass

    return moves_meta, ru_names

def parse_and_merge_json_files(json_paths):
    """Сбор и суммирование всех ТМ и TR дисков из JSON файлов."""
    merged_items = {}

    for path in json_paths:
        if not os.path.exists(path):
            continue
        with open(path, 'r', encoding='utf-8') as f:
            data = json.load(f)
            tms = data.get('tms', [])
            for item in tms:
                item_id = item.get('id')
                count = item.get('count', 1)
                name = item.get('name', '')
                item_type = item.get('type')

                if item_id not in merged_items:
                    merged_items[item_id] = {
                        'id': item_id,
                        'name': name,
                        'type': item_type,
                        'count': 0
                    }
                merged_items[item_id]['count'] += count
                if item_type and not merged_items[item_id].get('type'):
                    merged_items[item_id]['type'] = item_type

    tm_list = []
    tr_list = []

    for item_id, data in merged_items.items():
        if ':tr_' in item_id:
            tr_list.append(data)
        else:
            tm_list.append(data)

    return tm_list, tr_list

def extract_number(name_or_id):
    """Извлечение числового номера для правильной сортировки (1, 2, 10 вместо 1, 10, 2)."""
    match = re.search(r'(?:TM|TR)[-\s]*(\d+)', name_or_id, re.IGNORECASE)
    if match:
        return int(match.group(1))
    match_num = re.search(r'\d+', name_or_id)
    if match_num:
        return int(match_num.group(0))
    return 9999

def build_excel_catalog(tm_items, tr_items, moves_meta, ru_names, output_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )

    header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')

    zebra_fill = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
    white_fill = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')

    bold_font = Font(name='Segoe UI', size=10, bold=True)
    regular_font = Font(name='Segoe UI', size=10)

    sheets_data = [
        ('TM-disks', tm_items, 'TM'),
        ('TR-disks', tr_items, 'TR')
    ]

    for sheet_name, items, prefix in sheets_data:
        ws = wb.create_sheet(title=sheet_name)
        ws.views.sheetView[0].showGridLines = True

        # Заголовки
        headers = ['ID предмета', 'Номер', 'Название (EN / RU)', 'Тип атаки', 'Количество']
        ws.append(headers)

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = thin_border
        ws.row_dimensions[1].height = 28

        # Сортировка по номеру диска
        items_sorted = sorted(items, key=lambda x: extract_number(x['name'] or x['id']))

        total_count = 0

        for row_idx, item in enumerate(items_sorted, start=2):
            raw_id = item['id']
            full_name = item['name']
            count = item['count']
            total_count += count

            num_val = extract_number(full_name or raw_id)
            num_str = f"{prefix}-{num_val}"

            clean_move_key = raw_id.split(':')[-1].replace('tm_', '').replace('tr_', '').replace('hm_', '').lower()
            meta = moves_meta.get(clean_move_key, {})
            
            en_name = meta.get('formatname')
            if not en_name:
                if ':' in full_name:
                    en_name = full_name.split(':', 1)[1].strip()
                else:
                    en_name = clean_move_key.replace('_', ' ').title()

            ru_name = ru_names.get(clean_move_key, '')
            if ru_name and ru_name.lower() != en_name.lower():
                display_name = f"{en_name} ({ru_name})"
            else:
                display_name = en_name

            # Определение типа (из JSON тега или метаданных)
            typing = (item.get('type') or meta.get('typing', 'normal')).lower()
            type_info = TYPE_DATA.get(typing, {'ru': typing.capitalize(), 'bg': 'E2E8F0', 'fg': '000000'})
            type_display = f"{type_info['ru']} ({typing.capitalize()})"

            ws.cell(row=row_idx, column=1, value=raw_id)
            ws.cell(row=row_idx, column=2, value=num_str)
            ws.cell(row=row_idx, column=3, value=display_name)
            ws.cell(row=row_idx, column=4, value=type_display)
            ws.cell(row=row_idx, column=5, value=count)

            is_even = (row_idx % 2 == 0)
            row_fill = zebra_fill if is_even else white_fill

            ws.cell(row=row_idx, column=1).alignment = Alignment(horizontal='left', vertical='center')
            ws.cell(row=row_idx, column=2).alignment = Alignment(horizontal='center', vertical='center')
            ws.cell(row=row_idx, column=3).alignment = Alignment(horizontal='left', vertical='center')
            ws.cell(row=row_idx, column=4).alignment = Alignment(horizontal='center', vertical='center')
            ws.cell(row=row_idx, column=5).alignment = Alignment(horizontal='center', vertical='center')

            for col_idx in range(1, 6):
                c = ws.cell(row=row_idx, column=col_idx)
                c.border = thin_border
                c.font = regular_font
                if col_idx != 4:
                    c.fill = row_fill

            type_cell = ws.cell(row=row_idx, column=4)
            type_cell.fill = PatternFill(start_color=type_info['bg'], end_color=type_info['bg'], fill_type='solid')
            type_cell.font = Font(name='Segoe UI', size=10, bold=True, color=type_info['fg'])

            ws.cell(row=row_idx, column=2).font = bold_font
            ws.cell(row=row_idx, column=5).font = bold_font

            ws.row_dimensions[row_idx].height = 22

        total_row = len(items_sorted) + 2
        ws.cell(row=total_row, column=1, value="ИТОГО:")
        ws.cell(row=total_row, column=2, value=f"{len(items_sorted)} уник.")
        ws.cell(row=total_row, column=3, value="")
        ws.cell(row=total_row, column=4, value="")
        ws.cell(row=total_row, column=5, value=f"=SUM(E2:E{total_row-1})")

        ws.row_dimensions[total_row].height = 24
        total_fill = PatternFill(start_color='E2E8F0', end_color='E2E8F0', fill_type='solid')
        for col_idx in range(1, 6):
            c = ws.cell(row=total_row, column=col_idx)
            c.font = Font(name='Segoe UI', size=11, bold=True, color='0F172A')
            c.fill = total_fill
            c.border = thin_border
            if col_idx in [1, 2, 5]:
                c.alignment = Alignment(horizontal='center', vertical='center')

        ws.freeze_panes = 'A2'

        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or '')
                if val.startswith('='):
                    val = str(total_count)
                max_len = max(max_len, len(val))
            ws.column_dimensions[col_letter].width = max(max_len + 5, 14)

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    try:
        wb.save(output_path)
        print(f"Таблица успешно сохранена: {output_path}")
    except PermissionError:
        alt_path = output_path.replace('.xlsx', '_updated.xlsx')
        wb.save(alt_path)
        print(f"[ВНИМАНИЕ] Файл '{os.path.basename(output_path)}' открыт в Excel. Сохранено в: {alt_path}")

def main():
    base_dir = r'c:\Users\Foxi8\OneDrive\Рабочий стол\СозданиеПеревода\tasks\Disk_copy_barrel\docs'
    # Основной файл дампа (содержит данные всех просканированных бочек)
    main_json = os.path.join(base_dir, 'my_tms_inventory.json')
    json_paths = [main_json] if os.path.exists(main_json) else [
        os.path.join(base_dir, 'my_tms_inventory2.json'),
        os.path.join(base_dir, 'my_tms_inventory3.json'),
    ]

    print("1. Загрузка метаданных атак и русских названий...")
    moves_meta, ru_names = load_move_metadata()

    print("2. Объединение данных из 3-х бочек...")
    tm_items, tr_items = parse_and_merge_json_files(json_paths)
    print(f"   Найдено уникальных TM: {len(tm_items)}")
    print(f"   Найдено уникальных TR: {len(tr_items)}")

    output_xlsx = os.path.join(base_dir, 'Cobblemon_TM_TR_Inventory.xlsx')
    print("3. Генерация стилизованной Excel-таблицы...")
    build_excel_catalog(tm_items, tr_items, moves_meta, ru_names, output_xlsx)

if __name__ == '__main__':
    main()
